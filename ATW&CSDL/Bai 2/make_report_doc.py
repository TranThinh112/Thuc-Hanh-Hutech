import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def create_report():
    doc = Document()
    
    # Thiết lập lề 2cm (~0.79 inch)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Style chung
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Tiêu đề báo cáo
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("BÁO CÁO BÀI TẬP THỰC HÀNH\nXÂY DỰNG ỨNG DỤNG WEB FULLSTACK ĐƠN GIẢN")
    title_run.font.size = Pt(18)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF) # Blue 800

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle.add_run("Học phần: Phát triển ứng dụng Web | Đề bài: Bài 2 - Fullstack Login")
    sub_run.font.size = Pt(12)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    doc.add_paragraph() # Khoảng trống

    # 1. Mục tiêu
    h1 = doc.add_heading("1. Mục tiêu bài thực hành", level=1)
    p = doc.add_paragraph(
        "- Làm quen và thực hành với các công nghệ client và server cơ bản trong ứng dụng Web.\n"
        "- Hiểu sâu sắc cơ chế giao tiếp Frontend - Backend qua giao thức HTTP (AJAX/Fetch API).\n"
        "- Thực hành phát triển ứng dụng Web đa tầng (Multi-tier Fullstack), xác thực người dùng và quản lý phiên (Session/Cookie)."
    )

    # 2. Kiến trúc & Công nghệ
    doc.add_heading("2. Công nghệ sử dụng & Cấu trúc ứng dụng", level=1)
    p_tech = doc.add_paragraph()
    p_tech.add_run("• Frontend: ").bold = True
    p_tech.add_run("HTML5, CSS3 hiện đại (Glassmorphism, Flexbox, Responsive), Vanilla JavaScript (AJAX Fetch API, DOM manipulation).\n")
    p_tech.add_run("• Backend: ").bold = True
    p_tech.add_run("Node.js runtime, Express framework, Express-Session, Cookie-Parser.\n")
    p_tech.add_run("• Cơ chế xác thực: ").bold = True
    p_tech.add_run("RESTful API (POST /api/login, GET /api/me, POST /api/logout), kiểm tra hợp lệ 2 lớp (Client & Server), bảo mật trạng thái bằng HTTP-only Cookie.")

    # 3. Kết quả chạy & Hình ảnh minh họa
    doc.add_heading("3. Kết quả thử nghiệm & Hình ảnh minh chứng", level=1)

    steps = [
        {
            "num": "Hình 1",
            "title": "Giao diện Form đăng nhập ban đầu",
            "file": "screenshots/01_form_dang_nhap.png",
            "desc": "Giao diện được thiết kế hiện đại theo phong cách Glassmorphism với nền chuyển màu mượt mà. Form gồm 2 ô nhập liệu Username, Password, nút chuyển đổi ẩn/hiện mật khẩu và hộp thông tin gợi ý tài khoản."
        },
        {
            "num": "Hình 2",
            "title": "Kiểm tra dữ liệu nhập liệu rỗng (Client-side Validation)",
            "file": "screenshots/02_loi_de_trong.png",
            "desc": "Khi người dùng để trống và nhấn nút 'Đăng nhập', JavaScript lập tức chặn gửi form (e.preventDefault()), kích hoạt hiển thị cảnh báo đỏ trực quan dưới từng ô nhập liệu tương ứng."
        },
        {
            "num": "Hình 3",
            "title": "Xử lý khi đăng nhập thất bại (Sai tài khoản / mật khẩu)",
            "file": "screenshots/03_dang_nhap_that_bai.png",
            "desc": "Frontend gửi yêu cầu POST qua Fetch API lên server. Backend kiểm tra không khớp với tài khoản mặc định và trả về mã lỗi HTTP 401 Unauthorized. Giao diện hiển thị Alert Box màu đỏ mà không cần tải lại trang."
        },
        {
            "num": "Hình 4",
            "title": "Nhập đúng thông tin & Chức năng bật/tắt hiển thị mật khẩu",
            "file": "screenshots/04_nhap_dung_hien_mat_khau.png",
            "desc": "Người dùng nhập đúng tài khoản 'admin' và '12345'. Nút icon con mắt cho phép chuyển đổi loại input từ 'password' sang 'text' để người dùng kiểm tra lại mật khẩu trước khi gửi."
        },
        {
            "num": "Hình 5",
            "title": "Đăng nhập thành công & Giao diện Dashboard quản lý phiên",
            "file": "screenshots/05_dang_nhap_thanh_cong.png",
            "desc": "Server xác thực thành công trả về HTTP 200 kèm JSON thông tin người dùng và khởi tạo Session. Frontend tự động ẩn form đăng nhập, hiển thị thẻ Dashboard phiên làm việc với tên tài khoản, vai trò và thời gian đăng nhập."
        },
        {
            "num": "Hình 6",
            "title": "Duy trì phiên đăng nhập khi tải lại trang (Session Persistence)",
            "file": "screenshots/06_duy_tri_session_khi_f5.png",
            "desc": "Khi người dùng ấn F5 hoặc tải lại trình duyệt, trang web tự động gọi API GET /api/me. Server xác thực thông qua Session Cookie 'connect.sid' hợp lệ và duy trì trạng thái đăng nhập liên tục cho người dùng."
        }
    ]

    for item in steps:
        h2 = doc.add_heading(f"{item['num']}: {item['title']}", level=2)
        desc_p = doc.add_paragraph(item['desc'])
        desc_p.runs[0].font.size = Pt(10.5)

        img_path = item['file']
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(4)
            p_img.paragraph_format.space_after = Pt(8)
            run_img = p_img.add_run()
            run_img.add_picture(img_path, width=Inches(6.0))

            caption = doc.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption_run = caption.add_run(f"Minh chứng: {item['num']} - {item['title']}")
            caption_run.font.size = Pt(9.5)
            caption_run.font.italic = True
            caption_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            caption.paragraph_format.space_after = Pt(14)

    # 4. Đánh giá & Kết luận
    doc.add_heading("4. Kết luận & Đánh giá kết quả", level=1)
    p_concl = doc.add_paragraph(
        "Ứng dụng đã hoàn thành 100% các yêu cầu từ cơ bản đến nâng cao của bài tập:\n"
        "- Xây dựng đầy đủ Frontend và Backend, kết nối thông suốt qua giao thức HTTP (AJAX/JSON).\n"
        "- Trải nghiệm người dùng mượt mà, phản hồi tức thời không cần reload trang.\n"
        "- Cơ chế xác thực chặt chẽ, kiểm tra hợp lệ cả 2 phía và quản lý phiên an toàn bằng HTTP-Only Session Cookie."
    )

    doc.save("Bai 2.docx")
    print("Tao file Bai 2.docx thanh cong!")

create_report()
