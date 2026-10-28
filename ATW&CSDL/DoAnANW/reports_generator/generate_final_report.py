# -*- coding: utf-8 -*-
import sys
import os
import json
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

sys.stdout.reconfigure(encoding="utf-8")

SRC_PATH = r"C:/Users/Admin/Downloads/BAO_CAO_DO_AN_XAY_DUNG_WEBSITE_TMDT_PHONG_CHONG_GIAN_LAN_THANH_TOAN.docx"
OUT_DOWNLOAD = r"C:/Users/Admin/Downloads/BAO_CAO_DO_AN_XAY_DUNG_WEBSITE_TMDT_PHONG_CHONG_GIAN_LAN_THANH_TOAN_HOAN_CHINH.docx"
OUT_LOCAL = r"BAO_CAO_DO_AN_XAY_DUNG_WEBSITE_TMDT_PHONG_CHONG_GIAN_LAN_THANH_TOAN_HOAN_CHINH.docx"

print(f"Loading original document: {SRC_PATH}")
doc = docx.Document(SRC_PATH)
w_ns = nsdecls("w")

def format_run(run, font_name="Times New Roman", font_size=13, bold=False, italic=False, color=None):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color

def add_p(target_p, text="", bold=False, italic=False, size=13, space_before=0, space_after=6, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.3):
    p = target_p.insert_paragraph_before()
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        r = p.add_run(text)
        format_run(r, font_size=size, bold=bold, italic=italic)
    return p

def add_h1(target_p, text):
    p = target_p.insert_paragraph_before()
    if "Heading 1" in doc.styles:
        p.style = doc.styles["Heading 1"]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text)
    format_run(r, font_size=15, bold=True)
    return p

def add_h2(target_p, text):
    p = target_p.insert_paragraph_before()
    if "Heading 2" in doc.styles:
        p.style = doc.styles["Heading 2"]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    format_run(r, font_size=13.5, bold=True)
    return p

def add_h3(target_p, text):
    p = target_p.insert_paragraph_before()
    if "Heading 3" in doc.styles:
        p.style = doc.styles["Heading 3"]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    format_run(r, font_size=13, bold=True)
    return p

def add_bullet(target_p, title, desc):
    p = target_p.insert_paragraph_before()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.left_indent = Inches(0.25)
    r_bullet = p.add_run("•  ")
    format_run(r_bullet, font_size=13, bold=True)
    if title:
        r_title = p.add_run(title + ": ")
        format_run(r_title, font_size=13, bold=True)
    r_desc = p.add_run(desc)
    format_run(r_desc, font_size=13, bold=False)
    return p

def set_cell_background(cell, fill_hex):
    shd = parse_xml(f'<w:shd {w_ns} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def add_table_data(doc, target_p, headers, rows, col_widths, caption):
    if caption:
        p_cap = target_p.insert_paragraph_before()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(10)
        p_cap.paragraph_format.space_after = Pt(4)
        r_lbl = p_cap.add_run(caption)
        format_run(r_lbl, font_size=11, bold=True, italic=True)
        
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Header row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1F497D")
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        for r in p.runs:
            format_run(r, font_size=10.5, bold=True, color=RGBColor(255, 255, 255))
            
    # Data rows
    for r_idx, row_data in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        bg_col = "F8F9FA" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            if bg_col != "FFFFFF":
                set_cell_background(row_cells[c_idx], bg_col)
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            if c_idx == 0 or len(str(val)) < 12:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                format_run(r, font_size=10, bold=(c_idx == 0 and len(row_data) > 3))
                
    # Borders
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {w_ns}><w:top w:val="single" w:sz="6" w:space="0" w:color="1F497D"/><w:bottom w:val="single" w:sz="6" w:space="0" w:color="1F497D"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/><w:left w:val="none"/><w:right w:val="none"/><w:insideV w:val="none"/></w:tblBorders>')
    tblPr.append(borders)
    
    # Column widths
    for row in table.rows:
        for c_idx, w in enumerate(col_widths):
            row.cells[c_idx].width = Inches(w)
            
    target_p._element.addprevious(table._element)
    
    # Spacing after
    p_sp = target_p.insert_paragraph_before()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(6)

def add_figure_item(target_p, img_path, fig_id, fig_title, commentary):
    if not os.path.isabs(img_path):
        img_path = os.path.abspath(img_path)
    p_img = target_p.insert_paragraph_before()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(10)
    p_img.paragraph_format.space_after = Pt(4)
    r_img = p_img.add_run()
    r_img.add_picture(img_path, width=Inches(5.95))
    
    p_cap = target_p.insert_paragraph_before()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(8)
    r_lbl = p_cap.add_run(f"Hình {fig_id}: ")
    format_run(r_lbl, font_size=11, bold=True, italic=True)
    r_txt = p_cap.add_run(fig_title)
    format_run(r_txt, font_size=11, bold=False, italic=True)
    
    if commentary:
        for c in commentary:
            add_p(target_p, c, size=13, space_before=0, space_after=6, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

def add_toc_line(target_p, title, page_num, level=1):
    p = target_p.insert_paragraph_before()
    p.paragraph_format.line_spacing = 1.2
    if level == 1:
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Inches(0)
        font_size = 12
        bold = True
    elif level == 2:
        p.paragraph_format.space_before = Pt(1.5)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.left_indent = Inches(0.25)
        font_size = 11.5
        bold = False
    else: # level 3
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.left_indent = Inches(0.5)
        font_size = 11
        bold = False

    # Set right tab with dot leader at 8769 dxa (~6.09 in)
    pPr = p._p.get_or_add_pPr()
    tabs = parse_xml(f'<w:tabs {w_ns}><w:tab w:val="right" w:leader="dot" w:pos="8769"/></w:tabs>')
    pPr.append(tabs)

    r_title = p.add_run(title)
    format_run(r_title, font_size=font_size, bold=bold)
    r_tab = p.add_run(f"\t{page_num}")
    format_run(r_tab, font_size=font_size, bold=bold)
    return p


# ==============================================================================
# STEP 1: REPLACE CHAPTER 3
# ==============================================================================
print("Step 1: Locating and replacing Chapter 3...")
ch3_idx = None
ch4_idx = None
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt.startswith("CHƯƠNG 3:") and ch3_idx is None:
        ch3_idx = i
    elif txt.startswith("CHƯƠNG 4:") and ch4_idx is None:
        ch4_idx = i
        break

print(f"  Old Chapter 3 spans from paragraph {ch3_idx} to {ch4_idx-1}")
p_ch4 = doc.paragraphs[ch4_idx]

# Delete old paragraphs between ch3_idx and ch4_idx
old_paras = [doc.paragraphs[k] for k in range(ch3_idx, ch4_idx)]
for p in old_paras:
    p._element.getparent().remove(p._element)

# Now insert complete Chapter 3 before p_ch4
with open("chapter3_sec1.json", "r", encoding="utf-8") as f:
    sec1 = json.load(f)
with open("chapter3_sec2.json", "r", encoding="utf-8") as f:
    sec2 = json.load(f)
with open("chapter3_sec3.json", "r", encoding="utf-8") as f:
    sec3 = json.load(f)
with open("chapter3_sec4.json", "r", encoding="utf-8") as f:
    sec4 = json.load(f)
with open("chapter3_sec5.json", "r", encoding="utf-8") as f:
    sec5 = json.load(f)

print("  Inserting Chapter 3 Title and Section 3.1...")
add_h1(p_ch4, "CHƯƠNG 3: XÂY DỰNG WEBSITE THƯƠNG MẠI ĐIỆN TỬ VÀ THỰC NGHIỆM HỆ THỐNG PHÁT HIỆN, PHÒNG CHỐNG GIAN LẬN")

# SECTION 3.1
add_h2(p_ch4, sec1["title"])
add_p(p_ch4, sec1["intro"])
for sub in sec1["subsections"]:
    add_h3(p_ch4, sub["title"])
    for p_txt in sub.get("paragraphs", []):
        add_p(p_ch4, p_txt)
    for b_title, b_desc in sub.get("bullets", []):
        add_bullet(p_ch4, b_title, b_desc)

# SECTION 3.2
print("  Inserting Section 3.2...")
add_h2(p_ch4, sec2["title"])
add_p(p_ch4, sec2["intro"])
for sub in sec2["subsections"]:
    add_h3(p_ch4, sub["title"])
    for p_txt in sub.get("paragraphs", []):
        add_p(p_ch4, p_txt)
    for b_title, b_desc in sub.get("bullets", []):
        add_bullet(p_ch4, b_title, b_desc)
    if "table" in sub:
        t = sub["table"]
        add_table_data(doc, p_ch4, t["headers"], t["rows"], t["widths"], t["caption"])

# SECTION 3.3
print("  Inserting Section 3.3...")
add_h2(p_ch4, sec3["title"])
add_p(p_ch4, sec3["intro"])
for sub in sec3["subsections"]:
    add_h3(p_ch4, sub["title"])
    for p_txt in sub.get("paragraphs", []):
        add_p(p_ch4, p_txt)
    if "table_rules" in sub:
        t = sub["table_rules"]
        add_table_data(doc, p_ch4, t["headers"], t["rows"], t["widths"], t["caption"])
    for p_txt in sub.get("post_paragraphs", []):
        add_p(p_ch4, p_txt)
    if "table_decisions" in sub:
        t = sub["table_decisions"]
        add_table_data(doc, p_ch4, t["headers"], t["rows"], t["widths"], t["caption"])

# SECTION 3.4
print("  Inserting Section 3.4 with 14 UI screenshots...")
add_h2(p_ch4, sec4["title"])
add_p(p_ch4, sec4["intro"])
for sub in sec4["subsections"]:
    add_h3(p_ch4, sub["title"])
    if "intro" in sub:
        add_p(p_ch4, sub["intro"])
    for fig in sub.get("figures", []):
        print(f"    Embedding Hình {fig['id']}: {fig['title']}")
        add_figure_item(p_ch4, fig["file"], fig["id"], fig["title"], fig["commentary"])

# SECTION 3.5
print("  Inserting Section 3.5 with test scenarios and evaluation...")
add_h2(p_ch4, sec5["title"])
add_p(p_ch4, sec5["intro"])
for sub in sec5["subsections"]:
    add_h3(p_ch4, sub["title"])
    for p_txt in sub.get("paragraphs", []):
        add_p(p_ch4, p_txt)
    if "table" in sub:
        t = sub["table"]
        add_table_data(doc, p_ch4, t["headers"], t["rows"], t["widths"], t["caption"])
    if "table_confusion" in sub:
        t = sub["table_confusion"]
        add_table_data(doc, p_ch4, t["headers"], t["rows"], t["widths"], t["caption"])
    for p_txt in sub.get("post_paragraphs", []):
        add_p(p_ch4, p_txt)
    for p_txt in sub.get("post_conf_paragraphs", []):
        add_p(p_ch4, p_txt)
    if "table_metrics" in sub:
        t = sub["table_metrics"]
        add_table_data(doc, p_ch4, t["headers"], t["rows"], t["widths"], t["caption"])
    for p_txt in sub.get("final_paragraphs", []):
        add_p(p_ch4, p_txt)

# Page break before Chapter 4
p_pb_ch4 = p_ch4.insert_paragraph_before()
r_pb_ch4 = p_pb_ch4.add_run()
r_pb_ch4.add_break(WD_BREAK.PAGE)


# ==============================================================================
# STEP 2: POPULATE DANH MỤC HÌNH ẢNH
# ==============================================================================
print("Step 2: Populating Danh mục hình ảnh...")
dmha_idx = None
ch1_idx = None
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip().upper()
    if "DANH MỤC HÌNH ẢNH" in txt and dmha_idx is None:
        dmha_idx = i
    elif txt.startswith("CHƯƠNG 1:") and ch1_idx is None:
        ch1_idx = i
        break

p_dmha = doc.paragraphs[dmha_idx]
p_ch1 = doc.paragraphs[ch1_idx]

# Format DANH MỤC HÌNH ẢNH heading
p_dmha.text = "DANH MỤC HÌNH ẢNH"
p_dmha.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_dmha.paragraph_format.space_before = Pt(12)
p_dmha.paragraph_format.space_after = Pt(12)
for r in p_dmha.runs:
    format_run(r, font_size=15, bold=True)

print(f"  DANH MỤC HÌNH ẢNH at {dmha_idx}, CHƯƠNG 1 at {ch1_idx}")

# Remove empty paragraphs between DANH MỤC HÌNH ẢNH and CHƯƠNG 1
paras_to_remove = [doc.paragraphs[k] for k in range(dmha_idx + 1, ch1_idx)]
for p in paras_to_remove:
    p._element.getparent().remove(p._element)

FIGURE_PAGES = {
    "3.1": 46, "3.2": 47, "3.3": 47, "3.4": 48, "3.5": 49,
    "3.6": 49, "3.7": 50, "3.8": 50, "3.9": 51, "3.10": 52,
    "3.11": 52, "3.12": 53, "3.13": 53, "3.14": 54
}

for sub in sec4["subsections"]:
    for fig in sub["figures"]:
        p_fig = p_ch1.insert_paragraph_before()
        p_fig.paragraph_format.line_spacing = 1.25
        p_fig.paragraph_format.space_before = Pt(2)
        p_fig.paragraph_format.space_after = Pt(2)
        p_fig.paragraph_format.left_indent = Inches(0)
        
        # Right tab with dot leader
        pPr = p_fig._p.get_or_add_pPr()
        tabs = parse_xml(f'<w:tabs {w_ns}><w:tab w:val="right" w:leader="dot" w:pos="8769"/></w:tabs>')
        pPr.append(tabs)
        
        r_num = p_fig.add_run(f"Hình {fig['id']}: ")
        format_run(r_num, font_size=11.5, bold=True)
        r_title = p_fig.add_run(fig["title"])
        format_run(r_title, font_size=11.5, bold=False)
        r_tab = p_fig.add_run(f"\t{FIGURE_PAGES.get(fig['id'], 46)}")
        format_run(r_tab, font_size=11.5, bold=False)

# Add page break before Chapter 1
p_pb_ch1 = p_ch1.insert_paragraph_before()
r_pb_ch1 = p_pb_ch1.add_run()
r_pb_ch1.add_break(WD_BREAK.PAGE)


# ==============================================================================
# STEP 3: POPULATE MỤC LỤC
# ==============================================================================
print("Step 3: Populating Table of Contents (Mục lục)...")
mucluc_idx = None
dmha_new_idx = None
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip().upper()
    if txt == "MỤC LỤC" and mucluc_idx is None:
        mucluc_idx = i
    elif "DANH MỤC HÌNH ẢNH" in txt and dmha_new_idx is None:
        dmha_new_idx = i
        break

p_mucluc = doc.paragraphs[mucluc_idx]
p_dmha_target = doc.paragraphs[dmha_new_idx]

# Format MỤC LỤC heading
p_mucluc.text = "MỤC LỤC"
p_mucluc.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_mucluc.paragraph_format.space_before = Pt(12)
p_mucluc.paragraph_format.space_after = Pt(12)
for r in p_mucluc.runs:
    format_run(r, font_size=15, bold=True)

print(f"  MỤC LỤC at {mucluc_idx}, DANH MỤC HÌNH ẢNH at {dmha_new_idx}")

# Remove empty paragraphs between MỤC LỤC and DANH MỤC HÌNH ẢNH
paras_to_remove = [doc.paragraphs[k] for k in range(mucluc_idx + 1, dmha_new_idx)]
for p in paras_to_remove:
    p._element.getparent().remove(p._element)

TOC_DATA = [
    ("LỜI MỞ ĐẦU", 3, 1),
    ("LỜI CẢM ƠN", 4, 1),
    ("NHIỆM VỤ ĐỒ ÁN", 5, 1),
    ("MỤC LỤC", 8, 1),
    ("DANH MỤC HÌNH ẢNH", 10, 1),
    ("CHƯƠNG 1: TỔNG QUAN VỀ ĐỀ TÀI", 11, 1),
    ("1.1 Bối cảnh và lý do chọn đề tài", 11, 2),
    ("1.1.1 Sự phát triển của thương mại điện tử và thanh toán số hiện đại", 11, 3),
    ("1.1.2 Nguy cơ an ninh và tính cấp thiết của bài toán chống gian lận", 12, 3),
    ("1.2 Thực trạng an toàn thông tin trong thanh toán trực tuyến", 13, 2),
    ("1.2.1 Cơ chế trao đổi dữ liệu thanh toán trực tuyến", 13, 3),
    ("1.2.2 Các điểm mù bảo mật và rủi ro truyền thông dữ liệu", 14, 3),
    ("1.3 Giới thiệu đề tài nghiên cứu", 15, 2),
    ("1.4 Mục tiêu và phạm vi nghiên cứu của đồ án", 15, 2),
    ("1.4.1 Mục tiêu nghiên cứu", 15, 3),
    ("1.4.2 Phạm vi nghiên cứu", 16, 3),
    ("1.5 Phương pháp nghiên cứu", 16, 2),
    ("1.6 Bố cục của đồ án", 17, 2),
    ("CHƯƠNG 2: CƠ SỞ LÝ THUYẾT VÀ CÁC HÌNH THỨC GIAN LẬN THANH TOÁN", 18, 1),
    ("2.1 Bản chất và khái niệm gian lận thanh toán", 18, 2),
    ("2.1.1 Định nghĩa gian lận thanh toán trong TMĐT", 18, 3),
    ("2.1.2 Ba yếu tố cốt lõi của hành vi gian lận", 19, 3),
    ("2.1.3 Các đối tượng trong chuỗi giao dịch bị đe dọa", 20, 3),
    ("2.1.4 Động cơ và mục tiêu gian lận", 21, 3),
    ("2.2 Phân loại các hình thức gian lận thanh toán phổ biến", 22, 2),
    ("2.2.1 Gian lận từ phía người mua (Buyer Fraud)", 22, 3),
    ("2.2.2 Gian lận từ phía người bán (Merchant Fraud)", 23, 3),
    ("2.2.3 Các phương thức tấn công kỹ thuật số (Digital & Cyber Attacks)", 24, 3),
    ("2.2.4 Gian lận liên quan đến Ví điện tử và tiền kỹ thuật số", 25, 3),
    ("2.3 Phân tích tác động và hậu quả của gian lận thanh toán", 26, 2),
    ("2.4 Phân tích chuyên sâu các kịch bản gian lận thực tế", 27, 2),
    ("2.4.1 Kịch bản 1: Giả mạo trang thanh toán (Phishing Checkout Page)", 27, 3),
    ("2.4.2 Kịch bản 2: Thử thẻ tín dụng tự động bằng Bot (Card Testing Attack)", 28, 3),
    ("2.4.3 Kịch bản 3: Giao dịch di chuyển bất khả thi (Impossible Travel / Velocity Check)", 29, 3),
    ("2.5 Các giải pháp và cơ chế kỹ thuật phòng chống gian lận hiện đại", 30, 2),
    ("2.5.1 Bộ quy tắc kiểm tra giao dịch (Rule-Based Engine) và Blacklist/Whitelist", 30, 3),
    ("2.5.2 Kỹ thuật dấu vân tay thiết bị (Device Fingerprinting) và GeoIP", 30, 3),
    ("2.5.3 Cơ chế kiểm tra vận tốc giao dịch (Velocity Checking)", 31, 3),
    ("2.5.4 Xác thực khách hàng nâng cao (3D-Secure 2.0, OTP, MFA)", 31, 3),
    ("2.5.5 Tiêu chuẩn an toàn dữ liệu thanh toán PCI-DSS và Tokenization", 32, 3),
    ("2.5.6 Ứng dụng Trí tuệ nhân tạo trong chấm điểm rủi ro (Risk Scoring)", 32, 3),
    ("CHƯƠNG 3: XÂY DỰNG WEBSITE THƯƠNG MẠI ĐIỆN TỬ VÀ THỰC NGHIỆM HỆ THỐNG PHÁT HIỆN, PHÒNG CHỐNG GIAN LẬN", 33, 1),
    ("3.1 Khảo sát và phân tích yêu cầu hệ thống", 33, 2),
    ("3.1.1 Yêu cầu nghiệp vụ cho website thương mại điện tử", 33, 3),
    ("3.1.2 Yêu cầu đối với module phát hiện và phòng chống gian lận", 34, 3),
    ("3.1.3 Các yêu cầu phi chức năng", 35, 3),
    ("3.2 Thiết kế kiến trúc hệ thống", 36, 2),
    ("3.2.1 Kiến trúc tổng thể phần mềm và luồng tương tác các thành phần", 36, 3),
    ("3.2.2 Lựa chọn ngăn xếp công nghệ phát triển", 37, 3),
    ("3.2.3 Thiết kế cơ sở dữ liệu quan hệ", 38, 3),
    ("3.2.4 Thiết kế luồng xử lý giao dịch thanh toán kèm bộ lọc gian lận", 40, 3),
    ("3.3 Thiết kế và cài đặt các module phòng chống gian lận", 41, 2),
    ("3.3.1 Module kiểm tra vận tốc và vị trí địa lý bất thường (Impossible Travel)", 41, 3),
    ("3.3.2 Module phát hiện tấn công thử thẻ tự động và chống quét Bot (Card Testing)", 42, 3),
    ("3.3.3 Module kiểm tra dấu vân tay thiết bị và tính toàn vẹn phiên làm việc", 43, 3),
    ("3.3.4 Bộ máy chấm điểm rủi ro và ra quyết định giao dịch", 44, 3),
    ("3.4 Hiện thực hóa giao diện hệ thống", 46, 2),
    ("3.4.1 Giao diện phía người dùng", 46, 3),
    ("3.4.2 Giao diện quản trị hệ thống (Admin Dashboard)", 50, 3),
    ("3.5 Triển khai thử nghiệm và đánh giá thực nghiệm các kịch bản gian lận", 54, 2),
    ("3.5.1 Môi trường và công cụ triển khai thử nghiệm", 54, 3),
    ("3.5.2 Thực nghiệm Kịch bản 1: Kiểm thử phòng chống tấn công thử thẻ tự động (Card Testing Attack)", 55, 3),
    ("3.5.3 Thực nghiệm Kịch bản 2: Kiểm thử phòng chống gian lận di chuyển bất khả thi (Impossible Travel)", 56, 3),
    ("3.5.4 Thực nghiệm Kịch bản 3: Kiểm thử cơ chế xác thực bổ sung MFA khi phát hiện hành vi nghi vấn", 57, 3),
    ("3.5.5 Đánh giá hiệu năng và độ trễ của hệ thống", 58, 3),
    ("CHƯƠNG 4: KẾT LUẬN VÀ GIẢI PHÁP ĐỀ XUẤT", 60, 1),
    ("4.1 Tổng kết kết quả nghiên cứu và sản phẩm đồ án", 60, 2),
    ("4.2 Đánh giá ưu điểm và hạn chế của đề tài", 61, 2),
    ("4.2.1 Các ưu điểm nổi bật đã đạt được", 61, 3),
    ("4.2.2 Những hạn chế cần tiếp tục hoàn thiện", 62, 3),
    ("4.3 Giải pháp đề xuất toàn diện cho các bên liên quan", 63, 2),
    ("4.3.1 Đối với người tiêu dùng (Khách hàng)", 63, 3),
    ("4.3.2 Đối với doanh nghiệp kinh doanh TMĐT", 64, 3),
    ("4.3.3 Đối với ngân hàng, cổng thanh toán và cơ quan quản lý", 65, 3),
    ("4.4 Thông điệp kết luận về an toàn giao dịch số", 66, 2),
    ("4.5 Hướng phát triển trong tương lai", 66, 2),
    ("TÀI LIỆU THAM KHẢO", 68, 1),
]

for title, page, level in TOC_DATA:
    add_toc_line(p_dmha_target, title, page, level)

# Add page break before DANH MỤC HÌNH ẢNH
p_pb_dmha = p_dmha_target.insert_paragraph_before()
r_pb_dmha = p_pb_dmha.add_run()
r_pb_dmha.add_break(WD_BREAK.PAGE)

# Remove all footers across all sections as requested by user
print("Removing all footers across all sections...")
for sec in doc.sections:
    sec.footer.is_linked_to_previous = False
    for p in list(sec.footer.paragraphs):
        p.text = ""
        for child in list(p._element):
            if not child.tag.endswith("pPr"):
                p._element.remove(child)
    if hasattr(sec, "first_page_footer") and sec.first_page_footer:
        sec.first_page_footer.is_linked_to_previous = False
        for p in list(sec.first_page_footer.paragraphs):
            p.text = ""
            for child in list(p._element):
                if not child.tag.endswith("pPr"):
                    p._element.remove(child)
    if hasattr(sec, "even_page_footer") and sec.even_page_footer:
        sec.even_page_footer.is_linked_to_previous = False
        for p in list(sec.even_page_footer.paragraphs):
            p.text = ""
            for child in list(p._element):
                if not child.tag.endswith("pPr"):
                    p._element.remove(child)
    for child in list(sec._sectPr):
        if "footerReference" in child.tag:
            sec._sectPr.remove(child)

print("Saving final documents...")
print(f"  -> {OUT_LOCAL}")
doc.save(OUT_LOCAL)
print(f"  -> {OUT_DOWNLOAD}")
doc.save(OUT_DOWNLOAD)

print("SUCCESS: Complete report document successfully generated!")
