import os
import json
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = docx.Document()

# 1. Page setup: A4, 0.75 in margins
for s in doc.sections:
    s.page_width = Inches(8.27)
    s.page_height = Inches(11.69)
    s.top_margin = Inches(0.75)
    s.bottom_margin = Inches(0.75)
    s.left_margin = Inches(0.75)
    s.right_margin = Inches(0.75)

# Colors
COLOR_PRIMARY = RGBColor(26, 82, 118)     # #1A5276 Dark Blue
COLOR_SECONDARY = RGBColor(41, 128, 185) # #2980B9 Soft Blue
COLOR_DARK = RGBColor(44, 62, 80)        # #2C3E50 Text color
COLOR_MUTED = RGBColor(127, 140, 141)    # #7F8C8D Muted grey
COLOR_SUCCESS = RGBColor(39, 174, 96)    # #27AE60 Green
COLOR_DANGER = RGBColor(192, 57, 43)     # #C0392B Red

def set_run(run, font_name="Calibri", size_pt=10.5, bold=False, italic=False, color=COLOR_DARK):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color

def add_h1(text):
    h = doc.add_paragraph()
    h.paragraph_format.keep_with_next = True
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(4)
    run = h.add_run(text)
    set_run(run, "Calibri", 14, bold=True, color=COLOR_PRIMARY)
    return h

def add_h2(text):
    h = doc.add_paragraph()
    h.paragraph_format.keep_with_next = True
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(3)
    run = h.add_run(text)
    set_run(run, "Calibri", 12, bold=True, color=COLOR_SECONDARY)
    return h

def add_p(text="", bold_pre="", space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_pre:
        r_pre = p.add_run(bold_pre)
        set_run(r_pre, "Calibri", 10.5, bold=True, color=COLOR_DARK)
    if text:
        r_text = p.add_run(text)
        set_run(r_text, "Calibri", 10.5, bold=False, color=COLOR_DARK)
    return p

def add_bullet(bold_pre="", text=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.15
    if bold_pre:
        r_pre = p.add_run(bold_pre)
        set_run(r_pre, "Calibri", 10.5, bold=True, color=COLOR_DARK)
    if text:
        r_text = p.add_run(text)
        set_run(r_text, "Calibri", 10.5, bold=False, color=COLOR_DARK)
    return p

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_borders(cell, top="none", bottom="none", left="none", right="none", 
                     color="CCCCCC", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="{top}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{left}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{bottom}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{right}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

def add_image_with_caption(img_path, caption, width_in=6.5):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(2)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(width_in))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(8)
        run_cap = p_cap.add_run(caption)
        set_run(run_cap, "Calibri", 9, italic=True, color=COLOR_MUTED)

# ==================== HEADER BANNER ====================
title_p = doc.add_paragraph()
title_p.paragraph_format.space_before = Pt(0)
title_p.paragraph_format.space_after = Pt(2)
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_sub = title_p.add_run("BÁO CÁO BÀI TẬP THỰC HÀNH\n")
set_run(r_sub, "Calibri", 12, bold=True, color=COLOR_SECONDARY)
r_title = title_p.add_run("PHÂN TÍCH GIAO TIẾP HTTP VÀ KIẾN TRÚC WEB")
set_run(r_title, "Calibri", 18, bold=True, color=COLOR_PRIMARY)

# Info Table / Header Box
info_tbl = doc.add_table(rows=2, cols=2)
info_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
col_widths = [Inches(3.38), Inches(3.38)]
for row in info_tbl.rows:
    for i, w in enumerate(col_widths):
        row.cells[i].width = w

set_cell_background(info_tbl.cell(0, 0), "F2F4F4")
set_cell_background(info_tbl.cell(0, 1), "F2F4F4")
set_cell_background(info_tbl.cell(1, 0), "F2F4F4")
set_cell_background(info_tbl.cell(1, 1), "F2F4F4")

p = info_tbl.cell(0, 0).paragraphs[0]
p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(2)
r = p.add_run("Học phần: "); set_run(r, "Calibri", 9.5, bold=True)
r = p.add_run("An toàn Ứng dụng Web / Mạng máy tính"); set_run(r, "Calibri", 9.5)

p = info_tbl.cell(0, 1).paragraphs[0]
p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(2)
r = p.add_run("Công cụ sử dụng: "); set_run(r, "Calibri", 9.5, bold=True)
r = p.add_run("Chrome DevTools, Chromium Engine"); set_run(r, "Calibri", 9.5)

p = info_tbl.cell(1, 0).paragraphs[0]
p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(2)
r = p.add_run("Trang web phân tích: "); set_run(r, "Calibri", 9.5, bold=True)
r = p.add_run("The Internet (Heroku Login System)"); set_run(r, "Calibri", 9.5)

p = info_tbl.cell(1, 1).paragraphs[0]
p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(2)
r = p.add_run("Môi trường: "); set_run(r, "Calibri", 9.5, bold=True)
r = p.add_run("Windows 10/11 x64, HTTP/1.1 & HTTP/2"); set_run(r, "Calibri", 9.5)

p_div = doc.add_paragraph()
p_div.paragraph_format.space_before = Pt(6)
p_div.paragraph_format.space_after = Pt(6)

# ==================== PHẦN 1 ====================
add_h1("PHẦN 1: QUY TRÌNH THỰC HIỆN VÀ THU THẬP DỮ LIỆU BẰNG DEVELOPER TOOLS")
add_p("Bài thực hành thực hiện phân tích luồng trao đổi dữ liệu HTTP/HTTPS giữa trình duyệt web (Client) và máy chủ (Backend) trên hệ thống xác thực người dùng thực tế: ")
p_link = add_p()
r = p_link.add_run("URL mục tiêu: https://the-internet.herokuapp.com/login")
set_run(r, "Calibri", 10.5, bold=True, color=COLOR_SECONDARY)

add_h2("1. Các bước tiến hành trên Developer Tools:")
add_bullet("Bước 1: ", "Khởi chạy trình duyệt Google Chrome. Nhấn phím F12 (hoặc Ctrl + Shift + I) để mở bảng điều khiển Chrome Developer Tools.")
add_bullet("Bước 2: ", "Chuyển sang tab Network. Đảm bảo nút ghi nhận mạng (đèn đỏ Record Network Log) đang bật và tick chọn checkbox 'Preserve log' nhằm giữ lại toàn bộ luồng request trong quá trình chuyển hướng trang (Redirect).")
add_bullet("Bước 3: ", "Nhập URL https://the-internet.herokuapp.com/login vào thanh địa chỉ và nhấn Enter (hoặc F5 Reload). Quan sát danh sách các request tải mã nguồn trang chính, stylesheet CSS, JavaScript libraries và fonts chữ.")
add_bullet("Bước 4: ", "Nhập thông tin tài khoản hợp lệ (Username: tomsmith, Password: SuperSecretPassword!) vào biểu mẫu và nhấn nút 'Login'.")
add_bullet("Bước 5: ", "Dừng ghi nhận và tiến hành lọc, phân tích chi tiết từng request GET và POST đã được thực thi.")

add_image_with_caption("captures/screenshot_network_overview.png", "Hình 1: Tổng quan bảng điều khiển Chrome DevTools - Tab Network ghi nhận chuỗi request khi tải trang và đăng nhập")

# ==================== PHẦN 2 ====================
add_h1("PHẦN 2: PHÂN TÍCH CHI TIẾT CÁC REQUEST HTTP (GET VÀ POST)")
add_p("Trong quá trình tương tác, trình duyệt đã thực hiện gửi và nhận nhiều loại tài nguyên khác nhau. Dưới đây là phân tích chi tiết 2 request quan trọng nhất đại diện cho phương thức truy vấn hiển thị (GET) và gửi nhận dữ liệu xác thực (POST).")

add_h2("2.1. Phân tích Request loại GET (Tải giao diện trang đăng nhập)")
add_p("Request GET được trình duyệt tự động tạo và gửi đi ngay khi người dùng nhập URL hoặc tải lại trang web để lấy tài nguyên tài liệu HTML.")

add_image_with_caption("captures/screenshot_get_headers.png", "Hình 2: Chi tiết Headers, URL và Response Code của Request GET /login trên DevTools")

# Table of GET Details
tbl_get = doc.add_table(rows=6, cols=2)
tbl_get.alignment = WD_TABLE_ALIGNMENT.CENTER
for row in tbl_get.rows:
    row.cells[0].width = Inches(2.2)
    row.cells[1].width = Inches(4.57)

get_rows_data = [
    ("Thông tin phân tích", "Giá trị ghi nhận thực tế"),
    ("Request URL & Method", "GET https://the-internet.herokuapp.com/login"),
    ("Mã trạng thái (Status Code)", "200 OK (Thành công - Máy chủ phản hồi nội dung tài liệu HTML)"),
    ("Loại tài nguyên (Resource Type)", "document (MIME-Type: text/html; charset=utf-8) | Dung lượng: 2.5 KB"),
    ("Key Response Headers", "• Server: Heroku\n• Content-Type: text/html;charset=utf-8\n• X-Frame-Options: SAMEORIGIN (Chống Clickjacking)\n• X-XSS-Protection: 1; mode=block\n• X-Content-Type-Options: nosniff (Chống MIME-sniffing)"),
    ("Key Request Headers", "• Accept: text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,*/*;q=0.8\n• User-Agent: Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/153.0\n• Upgrade-Insecure-Requests: 1 (Yêu cầu nâng cấp HTTPS)")
]

for idx, (c1, c2) in enumerate(get_rows_data):
    cell1 = tbl_get.cell(idx, 0)
    cell2 = tbl_get.cell(idx, 1)
    if idx == 0:
        set_cell_background(cell1, "1A5276")
        set_cell_background(cell2, "1A5276")
        p1 = cell1.paragraphs[0]; r1 = p1.add_run(c1); set_run(r1, "Calibri", 10, bold=True, color=RGBColor(255,255,255))
        p2 = cell2.paragraphs[0]; r2 = p2.add_run(c2); set_run(r2, "Calibri", 10, bold=True, color=RGBColor(255,255,255))
    else:
        bg = "F9EBEA" if idx % 2 == 1 else "FFFFFF"
        set_cell_background(cell1, "F2F4F4")
        set_cell_background(cell2, bg)
        p1 = cell1.paragraphs[0]; r1 = p1.add_run(c1); set_run(r1, "Calibri", 9.5, bold=True, color=COLOR_DARK)
        p2 = cell2.paragraphs[0]; r2 = p2.add_run(c2); set_run(r2, "Calibri", 9.5, color=COLOR_DARK)
    set_cell_borders(cell1, top="single", bottom="single", left="single", right="single", color="BDC3C7")
    set_cell_borders(cell2, top="single", bottom="single", left="single", right="single", color="BDC3C7")

p_sp = doc.add_paragraph(); p_sp.paragraph_format.space_before = Pt(4); p_sp.paragraph_format.space_after = Pt(4)

add_h2("2.2. Phân tích Request loại POST (Xác thực thông tin tài khoản)")
add_p("Request POST được kích hoạt khi người dùng bấm nút 'Login' để gửi thông tin đăng nhập từ biểu mẫu (Form) lên máy chủ xử lý.")

add_image_with_caption("captures/screenshot_post_headers.png", "Hình 3: Chi tiết Request POST /authenticate, Form Payload và Redirect 303 See Other")

# Table of POST Details
tbl_post = doc.add_table(rows=7, cols=2)
tbl_post.alignment = WD_TABLE_ALIGNMENT.CENTER
for row in tbl_post.rows:
    row.cells[0].width = Inches(2.2)
    row.cells[1].width = Inches(4.57)

post_rows_data = [
    ("Thông tin phân tích", "Giá trị ghi nhận thực tế"),
    ("Request URL & Method", "POST https://the-internet.herokuapp.com/authenticate"),
    ("Mã trạng thái (Status Code)", "303 See Other (Chuyển hướng thành công sau khi xác thực hợp lệ)"),
    ("Dữ liệu Payload (Body)", "username=tomsmith&password=SuperSecretPassword%21\n(Mã hóa dạng URL-encoded)"),
    ("Header điều hướng (Location)", "Location: https://the-internet.herokuapp.com/secure (Chuyển tiếp sang trang quản trị)"),
    ("Key Response Headers", "• Server: Heroku\n• Content-Length: 0\n• Via: 2.0 heroku-router\n• Location: https://the-internet.herokuapp.com/secure"),
    ("Key Request Headers", "• Content-Type: application/x-www-form-urlencoded\n• Origin: https://the-internet.herokuapp.com\n• Referer: https://the-internet.herokuapp.com/login\n• User-Agent: Chrome/153.0 (Client Web Browser)")
]

for idx, (c1, c2) in enumerate(post_rows_data):
    cell1 = tbl_post.cell(idx, 0)
    cell2 = tbl_post.cell(idx, 1)
    if idx == 0:
        set_cell_background(cell1, "27AE60")
        set_cell_background(cell2, "27AE60")
        p1 = cell1.paragraphs[0]; r1 = p1.add_run(c1); set_run(r1, "Calibri", 10, bold=True, color=RGBColor(255,255,255))
        p2 = cell2.paragraphs[0]; r2 = p2.add_run(c2); set_run(r2, "Calibri", 10, bold=True, color=RGBColor(255,255,255))
    else:
        bg = "EAFAF1" if idx % 2 == 1 else "FFFFFF"
        set_cell_background(cell1, "F2F4F4")
        set_cell_background(cell2, bg)
        p1 = cell1.paragraphs[0]; r1 = p1.add_run(c1); set_run(r1, "Calibri", 9.5, bold=True, color=COLOR_DARK)
        p2 = cell2.paragraphs[0]; r2 = p2.add_run(c2); set_run(r2, "Calibri", 9.5, color=COLOR_DARK)
    set_cell_borders(cell1, top="single", bottom="single", left="single", right="single", color="BDC3C7")
    set_cell_borders(cell2, top="single", bottom="single", left="single", right="single", color="BDC3C7")

p_sp2 = doc.add_paragraph(); p_sp2.paragraph_format.space_before = Pt(4); p_sp2.paragraph_format.space_after = Pt(4)

add_h2("2.3. Bảng phân loại các loại tài nguyên web được tải trong phiên làm việc:")
add_p("Trong một phiên truy cập trang web, trình duyệt không chỉ tải văn bản HTML đơn thuần mà phải tải rất nhiều thành phần đa phương tiện và kịch bản khác nhau để hoàn thiện giao diện (Rendering Pipeline):")

tbl_res = doc.add_table(rows=6, cols=4)
tbl_res.alignment = WD_TABLE_ALIGNMENT.CENTER
col_w = [Inches(1.5), Inches(1.5), Inches(1.2), Inches(2.57)]
for row in tbl_res.rows:
    for i, w in enumerate(col_w):
        row.cells[i].width = w

res_data = [
    ("Loại tài nguyên", "MIME Type tiêu biểu", "Mã trạng thái", "Vai trò và ví dụ thực tế"),
    ("HTML (Document)", "text/html", "200 OK, 303", "Cấu trúc DOM trang web (login, secure)"),
    ("CSS (Stylesheet)", "text/css", "200 OK", "Định dạng giao diện, màu sắc, bố cục (app.css, sauce.css)"),
    ("JavaScript (Script)", "application/javascript", "200 OK", "Tương tác logic client-side (jquery, foundation.js)"),
    ("Image (Ảnh/Vector)", "image/svg+xml, png", "200 OK", "Biểu tượng đồ họa, icon, logo (forkme_right_sauce.svg)"),
    ("Fonts (Phông chữ)", "font/woff2", "200 OK", "Phông chữ hiển thị văn bản (atkinson, fraunces.woff2)")
]

for idx, (c1, c2, c3, c4) in enumerate(res_data):
    row = tbl_res.rows[idx]
    if idx == 0:
        for i in range(4):
            set_cell_background(row.cells[i], "1A5276")
            set_cell_borders(row.cells[i], "single", "single", "single", "single", "BDC3C7")
            p = row.cells[i].paragraphs[0]
            r = p.add_run([c1, c2, c3, c4][i]); set_run(r, "Calibri", 9.5, bold=True, color=RGBColor(255,255,255))
    else:
        bg = "F4F6F6" if idx % 2 == 1 else "FFFFFF"
        for i in range(4):
            set_cell_background(row.cells[i], bg)
            set_cell_borders(row.cells[i], "single", "single", "single", "single", "BDC3C7")
            p = row.cells[i].paragraphs[0]
            bold_val = True if i == 0 else False
            r = p.add_run([c1, c2, c3, c4][i]); set_run(r, "Calibri", 9.0, bold=bold_val, color=COLOR_DARK)

# ==================== PHẦN 3 ====================
add_h1("PHẦN 3: SƠ ĐỒ VÀ MÔ TẢ KIẾN TRÚC MÔ HÌNH 3 LỚP (3-TIER ARCHITECTURE)")
add_p("Hệ thống ứng dụng web được phân tách thành 3 tầng độc lập theo chuẩn mô hình kinh điển nhằm đảm bảo tính module hóa, khả năng bảo trì và nâng cao năng lực bảo mật:")

add_image_with_caption("captures/architecture_diagram.png", "Hình 4: Sơ đồ kiến trúc mô hình 3 lớp (Client - Server - Database) của hệ thống xác thực web")

add_h2("Mô tả chi tiết chức năng và thành phần của từng tầng:")

add_bullet("1. Tầng Trình diễn (Presentation Tier / Client): ", 
           "Bao gồm môi trường trình duyệt web (Google Chrome/Firefox) của người dùng cuối. Tầng này chịu trách nhiệm gửi các yêu cầu HTTP/HTTPS lên máy chủ, nhận dữ liệu HTML/CSS/JS, biên dịch cây DOM và kết xuất giao diện trực quan. Phía Client cũng thực hiện kiểm tra sơ bộ định dạng dữ liệu (Client-side validation) và lưu trữ thông tin phiên tạm thời (Session/Cookies).")

add_bullet("2. Tầng Ứng dụng & Xử lý nghiệp vụ (Application / Business Logic Tier): ", 
           "Bao gồm máy chủ Web Server (Heroku Router, Nginx) kết hợp với Backend Application Framework (Ruby on Rails / Sinatra / NodeJS). Tầng này đóng vai trò 'bộ não' của ứng dụng: tiếp nhận và giải mã HTTP Request, kiểm tra tính hợp lệ của tham số gửi lên, thực hiện quy trình xác thực nghiệp vụ (so khớp hash mật khẩu với cơ sở dữ liệu), khởi tạo Session ID và phản hồi HTTP Response tương ứng (200 OK, 303 Redirect hoặc 401 Unauthorized).")

add_bullet("3. Tầng Dữ liệu (Data Tier / Database): ", 
           "Bao gồm hệ quản trị cơ sở dữ liệu (PostgreSQL, MySQL hoặc SQLite). Tầng này có nhiệm vụ lưu trữ bền vững (Persistent Storage) thông tin người dùng, bảng phân quyền tài khoản, bảng nhật ký truy cập (Access Audit Log) và phiên hoạt động. Tầng ứng dụng giao tiếp với tầng dữ liệu thông qua các giao thức nội bộ bảo mật (Database Connection String / ORM / SQL Statements).")

# ==================== PHẦN 4 ====================
add_h1("PHẦN 4: XÁC ĐỊNH BỀ MẶT TẤN CÔNG (ATTACK SURFACE) VÀ GIẢI PHÁP PHÒNG CHỐNG")
add_p("Bề mặt tấn công (Attack Surface) là tổng thể tất cả các điểm vào (entry points) mà kẻ tấn công có thể khai thác nhằm xâm nhập trái phép, đánh cắp dữ liệu hoặc phá hoại tính toàn vẹn của ứng dụng. Dựa trên phân tích luồng HTTP và kiến trúc ở các phần trên, các bề mặt tấn công tiêu biểu gồm:")

add_h2("4.1. Bề mặt tấn công tại Biểu mẫu đăng nhập (Login Form / POST /authenticate)")
add_bullet("Kỹ thuật tấn công dò mật khẩu (Brute-Force & Credential Stuffing): ", 
           "Kẻ tấn công sử dụng công cụ tự động (như Burp Suite Intruder, Hydra) gửi hàng ngàn request POST chứa danh sách tài khoản/mật khẩu phổ biến nhằm chiếm quyền tài khoản.\n"
           "→ Giải pháp phòng ngừa: Áp dụng cơ chế Rate Limiting (giới hạn số lượt request/phút theo IP), triển khai CAPTCHA khi đăng nhập sai nhiều lần và kích hoạt xác thực hai yếu tố (2FA/MFA).")

add_bullet("Tấn công Tiêm mã SQL (SQL Injection - SQLi): ", 
           "Nếu trường 'username' hoặc 'password' được Backend ghép chuỗi trực tiếp vào câu lệnh truy vấn SQL (ví dụ: SELECT * FROM users WHERE username = '...' AND password = '...'), kẻ tấn công có thể chèn các chuỗi payload như ' OR '1'='1' -- để đăng nhập trái phép mà không cần mật khẩu.\n"
           "→ Giải pháp phòng ngừa: Luôn sử dụng Parameterized Queries (Prepared Statements) hoặc thư viện ORM, không bao giờ nối chuỗi dữ liệu đầu vào trực tiếp vào câu lệnh SQL.")

add_bullet("Tấn công Đánh cắp phiên qua trung gian (Man-in-the-Middle - MitM): ", 
           "Nếu đường truyền không mã hóa HTTPS hoặc cấu hình SSL yếu, kẻ tấn công trên cùng mạng Wi-Fi có thể bắt gói tin chứa plain-text username và password.\n"
           "→ Giải pháp phòng ngừa: Bắt buộc sử dụng giao thức HTTPS toàn trang, thiết lập Header HSTS (Strict-Transport-Security) để ép trình duyệt luôn kết nối qua kênh mã hóa an toàn.")

add_h2("4.2. Bề mặt tấn công tại Cơ chế quản lý Phiên (Session & Cookie Management)")
add_bullet("Tấn công Chiếm đoạt phiên làm việc (Session Hijacking): ", 
           "Nếu Session Token không được bảo vệ bằng các cờ an toàn, hacker có thể đánh cắp token qua mã độc XSS hoặc mạng công cộng để mạo danh người dùng hợp lệ.\n"
           "→ Giải pháp phòng ngừa: Thiết lập đầy đủ các cờ an toàn cho Cookie: HttpOnly (ngăn JavaScript đọc cookie), Secure (chỉ truyền qua HTTPS) và SameSite=Strict/Lax (chống tấn công CSRF).")

add_h2("4.3. Bề mặt tấn công tại các thành phần mở rộng (Search Box, File Upload, v.v.)")
add_bullet("Tấn công XSS (Cross-Site Scripting) qua ô tìm kiếm / hiển thị dữ liệu: ", 
           "Chèn mã kịch bản độc hại <script>...</script> vào các trường nhập liệu để thực thi mã độc trên trình duyệt nạn nhân.\n"
           "→ Giải pháp: Áp dụng Content Security Policy (CSP), mã hóa ngữ cảnh (Contextual Output Encoding) và kiểm tra lọc dữ liệu đầu vào (Input Sanitization).")

add_bullet("Tấn công Tải lên tệp độc hại (Malicious File Upload): ", 
           "Người dùng tải lên shell script (.php, .jsp, .exe) mạo danh tệp ảnh để giành quyền kiểm soát máy chủ Web.\n"
           "→ Giải pháp: Kiểm tra kỹ định dạng tệp (Magic Bytes), lưu trữ tệp tải lên ở máy chủ tách biệt (như AWS S3) và tắt hoàn toàn quyền thực thi mã nguồn trong thư mục chứa file.")

# ==================== PHẦN 5 ====================
add_h1("PHẦN 5: NHẬN XÉT CÁ NHÂN VÀ BÀI HỌC KINH NGHIỆM")
add_p("Sau khi hoàn thành bài thực hành phân tích giao tiếp HTTP và kiến trúc ứng dụng web, em rút ra các bài học và nhận xét cá nhân cốt lõi sau:")

add_bullet("Về giao tiếp HTTP: ", 
           "HTTP là giao thức không lưu trạng thái (Stateless), do đó việc duy trì đăng nhập của người dùng hoàn toàn phụ thuộc vào cơ chế Cookie và Session. Nhờ công cụ Developer Tools, luồng dữ liệu trở nên cực kỳ trực quan: từ việc tải HTML ban đầu (GET), gửi thông tin xác thực (POST), nhận tín hiệu điều hướng (303 See Other) cho đến khi nạp đầy đủ các tài nguyên phụ trợ (CSS, JS, Fonts).")

add_bullet("Về góc độ An toàn thông tin: ", 
           "Mọi dữ liệu gửi từ phía Client lên máy chủ (bao gồm Header, Body, Cookie, User-Agent) đều có thể bị chỉnh sửa, làm giả hoặc đánh cắp một cách dễ dàng bằng các công cụ chuyên dụng. Do đó, nguyên tắc 'Zero Trust' và 'Never Trust Client Input' phải luôn được áp dụng triệt để ở tầng Backend Logic. Kiểm tra dữ liệu ở phía Client chỉ có giá trị cải thiện trải nghiệm người dùng (UX), hoàn toàn không có giá trị bảo mật độc lập.")

add_bullet("Về vai trò của Security Headers: ", 
           "Việc cấu hình chính xác các Response Headers như X-Frame-Options, X-Content-Type-Options, HSTS và Content-Security-Policy đóng vai trò là một lớp phòng thủ chiều sâu (Defense-in-Depth) vô cùng quan trọng giúp bảo vệ ứng dụng trước các cuộc tấn công Clickjacking, MIME-confusion hay XSS.")

# Save Document
doc.save("Bao_Cao_Thuc_Hanh_HTTP_Va_Kien_Truc_Web.docx")
print("Report generated successfully as Bao_Cao_Thuc_Hanh_HTTP_Va_Kien_Truc_Web.docx")
