import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = docx.Document()

# Page setup: A4, 0.75 in margins (approx 1.9 cm)
sections = doc.sections
for s in sections:
    s.page_width = Inches(8.27)
    s.page_height = Inches(11.69)
    s.top_margin = Inches(0.75)
    s.bottom_margin = Inches(0.75)
    s.left_margin = Inches(0.75)
    s.right_margin = Inches(0.75)

# Color Palette
COLOR_PRIMARY = RGBColor(26, 82, 118)     # #1A5276 Dark Blue
COLOR_SECONDARY = RGBColor(41, 128, 185) # #2980B9 Soft Blue
COLOR_DARK = RGBColor(44, 62, 80)        # #2C3E50 Text color
COLOR_MUTED = RGBColor(127, 140, 141)    # #7F8C8D Muted grey
COLOR_SUCCESS = RGBColor(39, 174, 96)    # #27AE60 Green
COLOR_DANGER = RGBColor(192, 57, 43)     # #C0392B Red

def set_run_font(run, font_name="Calibri", size_pt=11, bold=False, italic=False, color=COLOR_DARK):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color

def add_header_styled(doc, text, level=1):
    h = doc.add_paragraph()
    h.paragraph_format.keep_with_next = True
    if level == 1:
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        run = h.add_run(text)
        set_run_font(run, "Calibri", 15, bold=True, color=COLOR_PRIMARY)
        # Add bottom border or background line
    elif level == 2:
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        run = h.add_run(text)
        set_run_font(run, "Calibri", 12.5, bold=True, color=COLOR_SECONDARY)
    elif level == 3:
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(2)
        run = h.add_run(text)
        set_run_font(run, "Calibri", 11, bold=True, color=COLOR_DARK)
    return h

def add_p(doc, text="", bold_prefix="", space_after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        set_run_font(r_pre, "Calibri", 10.5, bold=True, color=COLOR_DARK)
    if text:
        r_text = p.add_run(text)
        set_run_font(r_text, "Calibri", 10.5, bold=False, color=COLOR_DARK)
    return p

def add_bullet(doc, bold_prefix="", text=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        set_run_font(r_pre, "Calibri", 10.5, bold=True, color=COLOR_DARK)
    if text:
        r_text = p.add_run(text)
        set_run_font(r_text, "Calibri", 10.5, bold=False, color=COLOR_DARK)
    return p

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_callout(doc, title, content, border_color="2980B9", bg_color="EBF5FB"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.77)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    # Left border only
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run(title + "\n")
    set_run_font(r1, "Calibri", 10.5, bold=True, color=COLOR_PRIMARY)
    r2 = p.add_run(content)
    set_run_font(r2, "Calibri", 10, italic=False, color=COLOR_DARK)
    
    # Empty space after callout
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(4)

print("Base helper functions defined.")
