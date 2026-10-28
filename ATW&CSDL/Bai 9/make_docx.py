import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = Document()

# Thiet lap margin
sections = doc.sections
for section in sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Helper styling
def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

# Title
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_run = title_p.add_run("BÁO CÁO ĐÁNH GIÁ AN TOÀN THÔNG TIN ỨNG DỤNG WEB\n")
title_run.font.name = 'Arial'
title_run.font.size = Pt(18)
title_run.font.bold = True
title_run.font.color.rgb = RGBColor(30, 41, 59)

sub_run = title_p.add_run("BÀI THỰC HÀNH SỐ 9: ĐÁNH GIÁ VÀ BÁO CÁO LỖ HỔNG WEB THEO CHUẨN OWASP\n")
sub_run.font.name = 'Arial'
sub_run.font.size = Pt(13)
sub_run.font.bold = True
sub_run.font.color.rgb = RGBColor(37, 99, 235)

info_p = doc.add_paragraph()
info_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
info_run = info_p.add_run("Hệ thống mục tiêu: SecureApp Vulnerable Web Lab (http://localhost:3000)\nThời gian thực hiện: Tháng 10/2026 | Phương pháp: Black-box / Gray-box Testing\n")
info_run.font.name = 'Arial'
info_run.font.size = Pt(10)
info_run.font.italic = True
info_run.font.color.rgb = RGBColor(100, 116, 139)

doc.add_paragraph("―" * 55).alignment = WD_ALIGN_PARAGRAPH.CENTER

# 1. Mục tiêu và phạm vi
h1 = doc.add_heading("1. MỤC TIÊU VÀ PHẠM VI ĐÁNH GIÁ (SCOPE)", level=1)
h1.style.font.color.rgb = RGBColor(30, 41, 59)

p = doc.add_paragraph()
p.add_run("• Mục tiêu: ").bold = True
p.add_run("Thực hành quy trình rà soát, đánh giá an toàn thông tin trên ứng dụng web mẫu; phát hiện các lỗ hổng bảo mật phổ biến theo OWASP Top 10; chụp lại bằng chứng khai thác và đề xuất biện pháp khắc phục kỹ thuật.")

p = doc.add_paragraph()
p.add_run("• Công cụ sử dụng: ").bold = True
p.add_run("Burp Suite Community Edition, Trình duyệt Web (Chrome/Firefox/Playwright), Python PoC Automation.")

# Hinh anh 1: Trang chu
doc.add_paragraph("Giao diện tổng quan của Ứng dụng Web Lab mục tiêu:")
img_p = doc.add_paragraph()
img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
img_p.add_run().add_picture("C:/Users/Admin/Downloads/Bai 9/screenshots/01_trang_chu.png", width=Inches(6.0))
cap1 = doc.add_paragraph("Hình 1: Giao diện Trang chủ SecureApp Vulnerable Lab")
cap1.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap1.runs[0].font.size = Pt(9)
cap1.runs[0].font.italic = True

# 2. Bảng tổng hợp lỗ hổng
h2 = doc.add_heading("2. TỔNG HỢP KẾT QUẢ ĐÁNH GIÁ (EXECUTIVE SUMMARY)", level=1)
h2.style.font.color.rgb = RGBColor(30, 41, 59)

table = doc.add_table(rows=6, cols=5)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ["Mã lỗ hổng", "Tên lỗ hổng bảo mật", "Vị trí / Endpoint", "Mức độ rủi ro", "OWASP Category"]
data = [
    ["VULN-01", "SQL Injection (Auth Bypass & UNION)", "/login & /products", "CRITICAL (9.8)", "A03:2021 - Injection"],
    ["VULN-02", "Reflected Cross-Site Scripting (XSS)", "/products?q=", "HIGH (8.2)", "A03:2021 - Injection"],
    ["VULN-03", "Stored Cross-Site Scripting (XSS)", "/feedback", "HIGH (8.0)", "A03:2021 - Injection"],
    ["VULN-04", "Insecure Direct Object Reference (IDOR)", "/profile?id=", "HIGH (7.5)", "A01:2021 - Broken Access Control"],
    ["VULN-05", "Privilege Escalation via Header Bypass", "/admin", "CRITICAL (9.1)", "A01:2021 - Broken Access Control"]
]

# Style Table Header
hdr_cells = table.rows[0].cells
for idx, title in enumerate(headers):
    hdr_cells[idx].text = title
    set_cell_background(hdr_cells[idx], "1E293B")
    p = hdr_cells[idx].paragraphs[0]
    p.runs[0].font.name = 'Arial'
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
    p.runs[0].font.size = Pt(9.5)

# Style Data Rows
for r_idx, row in enumerate(data):
    row_cells = table.rows[r_idx + 1].cells
    for c_idx, val in enumerate(row):
        row_cells[c_idx].text = val
        p = row_cells[c_idx].paragraphs[0]
        p.runs[0].font.name = 'Arial'
        p.runs[0].font.size = Pt(9)
        if "CRITICAL" in val:
            p.runs[0].font.bold = True
            p.runs[0].font.color.rgb = RGBColor(185, 28, 28)
        elif "HIGH" in val:
            p.runs[0].font.bold = True
            p.runs[0].font.color.rgb = RGBColor(194, 65, 12)
        bg = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
        set_cell_background(row_cells[c_idx], bg)

doc.add_paragraph()

# 3. Chi tiết phát hiện và hình ảnh minh họa PoC
h3 = doc.add_heading("3. CHI TIẾT LỖ HỔNG VÀ BẰNG CHỨNG KHAI THÁC (PROOF OF CONCEPT)", level=1)
h3.style.font.color.rgb = RGBColor(30, 41, 59)

# 3.1 SQL Injection
doc.add_heading("3.1. VULN-01: SQL Injection (Bypass Authentication & Data Leak)", level=2)
p = doc.add_paragraph()
p.add_run("• Mô tả: ").bold = True
p.add_run("Ứng dụng thực hiện nối chuỗi đầu vào của người dùng trực tiếp vào câu lệnh SQL mà không sử dụng cơ chế Parameterized Query / Prepared Statements.\n")
p.add_run("• Khai thác Vượt qua Đăng nhập: ").bold = True
p.add_run("Tại form đăng nhập /login, truyền payload: ")
p.add_run("admin' --").bold = True
p.add_run(" vào trường username. Khi đó câu truy vấn thành: SELECT * FROM users WHERE username = 'admin' -- AND password = '...'. Đoạn kiểm tra mật khẩu bị comment bỏ qua hoàn toàn.")

# Hinh 2 + 3
doc.add_paragraph().alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.paragraphs[-1].add_run().add_picture("C:/Users/Admin/Downloads/Bai 9/screenshots/02_sqli_login_input.png", width=Inches(5.5))
cap2 = doc.add_paragraph("Hình 2: Chèn payload SQLi admin' -- vào form Đăng nhập")
cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER; cap2.runs[0].font.size = Pt(9); cap2.runs[0].font.italic = True

doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph().alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.paragraphs[-1].add_run().add_picture("C:/Users/Admin/Downloads/Bai 9/screenshots/03_sqli_login_success.png", width=Inches(5.5))
cap3 = doc.add_paragraph("Hình 3: Đăng nhập thành công vào tài khoản Quản trị viên (admin) mà không cần mật khẩu")
cap3.alignment = WD_ALIGN_PARAGRAPH.CENTER; cap3.runs[0].font.size = Pt(9); cap3.runs[0].font.italic = True

p_union = doc.add_paragraph()
p_union.add_run("• Khai thác Trích xuất Dữ liệu (UNION-based): ").bold = True
p_union.add_run("Tại ô tìm kiếm /products, chèn payload: ' UNION SELECT id, username, password, role FROM users --. Hệ thống trích xuất toàn bộ bảng users (bao gồm mật khẩu dạng thô) lên màn hình sản phẩm.")

doc.add_paragraph().alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.paragraphs[-1].add_run().add_picture("C:/Users/Admin/Downloads/Bai 9/screenshots/04_sqli_union_leak.png", width=Inches(5.5))
cap4 = doc.add_paragraph("Hình 4: Toàn bộ thông tin tài khoản và mật khẩu bị lộ do UNION SQL Injection")
cap4.alignment = WD_ALIGN_PARAGRAPH.CENTER; cap4.runs[0].font.size = Pt(9); cap4.runs[0].font.italic = True

p_fix1 = doc.add_paragraph()
p_fix1.add_run("✔ Giải pháp khắc phục: ").bold = True
p_fix1.add_run("Sử dụng Prepared Statements (Parameterized Queries) của SQLite / ORM. Tuyệt đối không nối chuỗi dữ liệu người dùng vào câu lệnh SQL.")

# 3.2 Reflected & Stored XSS
doc.add_heading("3.2. VULN-02 & VULN-03: Cross-Site Scripting (Reflected & Stored XSS)", level=2)
p = doc.add_paragraph()
p.add_run("• Mô tả: ").bold = True
p.add_run("Ứng dụng in trực tiếp dữ liệu do người dùng nhập vào giao diện HTML mà không qua bộ lọc mã hóa thực thể (HTML Entity Encoding).\n")
p.add_run("• Reflected XSS: ").bold = True
p.add_run("Tại trang tìm kiếm /products?q=<script>alert('Reflected-XSS')</script>, script lập tức được thực thi trong phiên làm việc của nạn nhân khi mở liên kết.\n")
p.add_run("• Stored XSS: ").bold = True
p.add_run("Tại form góp ý /feedback, kẻ tấn công gửi payload: <script>alert(document.cookie)</script>. Mã độc được lưu cố định vào cơ sở dữ liệu và tự động kích hoạt mỗi khi người dùng khác truy cập.")

doc.add_paragraph().alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.paragraphs[-1].add_run().add_picture("C:/Users/Admin/Downloads/Bai 9/screenshots/05_reflected_xss.png", width=Inches(5.5))
cap5 = doc.add_paragraph("Hình 5: Khai thác thành công Reflected XSS trên chức năng tìm kiếm sản phẩm")
cap5.alignment = WD_ALIGN_PARAGRAPH.CENTER; cap5.runs[0].font.size = Pt(9); cap5.runs[0].font.italic = True

doc.add_paragraph().alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.paragraphs[-1].add_run().add_picture("C:/Users/Admin/Downloads/Bai 9/screenshots/06_stored_xss_list.png", width=Inches(5.5))
cap6 = doc.add_paragraph("Hình 6: Stored XSS được lưu vào hệ thống và hiển thị trong danh sách Góp ý")
cap6.alignment = WD_ALIGN_PARAGRAPH.CENTER; cap6.runs[0].font.size = Pt(9); cap6.runs[0].font.italic = True

p_fix2 = doc.add_paragraph()
p_fix2.add_run("✔ Giải pháp khắc phục: ").bold = True
p_fix2.add_run("Mã hóa toàn bộ đầu ra (Output Encoding) bằng các hàm encodeHTML; Thiết lập Content Security Policy (CSP); Bổ sung cờ HttpOnly cho Cookie để chống đánh cắp phiên.")

# 3.3 IDOR
doc.add_heading("3.3. VULN-04: Insecure Direct Object Reference (IDOR / BOLA)", level=2)
p = doc.add_paragraph()
p.add_run("• Mô tả: ").bold = True
p.add_run("Tại chức năng xem hồ sơ /profile?id=X, máy chủ chỉ nhận tham số id từ URL và truy vấn trả về dữ liệu mà không hề xác thực người gửi yêu cầu có phải là chủ sở hữu tài nguyên hay không.")

doc.add_paragraph().alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.paragraphs[-1].add_run().add_picture("C:/Users/Admin/Downloads/Bai 9/screenshots/07_idor_profile_id2.png", width=Inches(5.5))
cap7 = doc.add_paragraph("Hình 7: Đọc trộm hồ sơ người dùng John Doe bằng cách thay đổi ?id=2")
cap7.alignment = WD_ALIGN_PARAGRAPH.CENTER; cap7.runs[0].font.size = Pt(9); cap7.runs[0].font.italic = True

p_fix3 = doc.add_paragraph()
p_fix3.add_run("✔ Giải pháp khắc phục: ").bold = True
p_fix3.add_run("Kiểm tra quyền sở hữu (Access Control Check) tại tầng backend trước khi trả về dữ liệu: req.user.id === requestedId.")

# 3.4 Privilege Escalation
doc.add_heading("3.4. VULN-05: Broken Access Control & Privilege Escalation trên /admin", level=2)
p = doc.add_paragraph()
p.add_run("• Mô tả: ").bold = True
p.add_run("Trang quản trị /admin chặn người dùng thông thường bằng mã 403 Forbidden. Tuy nhiên, mã nguồn lại tin cậy HTTP Header X-Role-Override: admin do client gửi lên.")

doc.add_paragraph().alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.paragraphs[-1].add_run().add_picture("C:/Users/Admin/Downloads/Bai 9/screenshots/08_admin_403_forbidden.png", width=Inches(5.5))
cap8 = doc.add_paragraph("Hình 8: Truy cập thông thường vào /admin bị từ chối (403 Forbidden)")
cap8.alignment = WD_ALIGN_PARAGRAPH.CENTER; cap8.runs[0].font.size = Pt(9); cap8.runs[0].font.italic = True

doc.add_paragraph().alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.paragraphs[-1].add_run().add_picture("C:/Users/Admin/Downloads/Bai 9/screenshots/09_admin_bypass_header.png", width=Inches(5.5))
cap9 = doc.add_paragraph("Hình 9: Leo thang đặc quyền Admin thành công khi chèn Header X-Role-Override: admin")
cap9.alignment = WD_ALIGN_PARAGRAPH.CENTER; cap9.runs[0].font.size = Pt(9); cap9.runs[0].font.italic = True

p_fix4 = doc.add_paragraph()
p_fix4.add_run("✔ Giải pháp khắc phục: ").bold = True
p_fix4.add_run("Không bao giờ tin tưởng các HTTP Header do client tự gửi lên để phân quyền. Phân quyền người dùng bắt buộc phải được xác minh từ Session/Token phía Server.")

# 4. Kết luận
doc.add_heading("4. KẾT LUẬN VÀ BÀI HỌC KINH NGHIỆM", level=1)
doc.add_paragraph("Bài thực hành đã giúp nắm vững quy trình kiểm thử đánh giá an toàn thông tin theo chuẩn công nghiệp:\n"
                  "1. Thu thập thông tin và lập bản đồ các chức năng trọng yếu (Authentication, Search, User Input, Admin Control).\n"
                  "2. Sử dụng công cụ Proxy (Burp Suite) để can thiệp, chỉnh sửa tham số và gửi payload tấn công.\n"
                  "3. Xác định bản chất lỗ hổng từ mã nguồn và đề xuất biện pháp phòng vệ theo chiều sâu (Defense-in-depth).")

# Luu file
output_path = "C:/Users/Admin/Downloads/Bai 9/Bai 9.docx"
doc.save(output_path)
print(f"Bao cao Word da duoc luu tai: {output_path}")
